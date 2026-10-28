import os
import re
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class ExportService:
    @staticmethod
    def generate_lit_review_docx(review_data: Dict[str, Any], output_path: str) -> str:
        """
        Creates an academic-formatted DOCX file for a literature review.
        """
        doc = Document()

        # Set page margins (1 inch academic standard)
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(review_data.get("title", "Literature Review"))
        title_run.font.name = "Times New Roman"
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)

        # Subtitle / Metadata
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_p.add_run(
            f"Domain: {review_data.get('domain', 'Academic Research')} | "
            f"Style: {review_data.get('citation_style', 'APA 7th')} | "
            f"Target Journal: {review_data.get('target_journal', 'Peer-Reviewed Journal')}"
        )
        meta_run.font.name = "Times New Roman"
        meta_run.font.size = Pt(10)
        meta_run.font.italic = True
        meta_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph()  # Spacing

        # Disclaimer
        disc_p = doc.add_paragraph()
        disc_run = disc_p.add_run(f"Disclaimer: {review_data.get('disclaimer', '')}")
        disc_run.font.name = "Times New Roman"
        disc_run.font.size = Pt(9)
        disc_run.font.italic = True
        disc_run.font.color.rgb = RGBColor(148, 163, 184)

        doc.add_paragraph()

        # Thematic Sections
        for sec in review_data.get("sections", []):
            h = doc.add_heading(sec.get("title", "Thematic Section"), level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)

            paragraphs = sec.get("content", "").split("\n\n")
            for p_text in paragraphs:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.first_line_indent = Inches(0.5)
                run = p.add_run(p_text.strip())
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

        # Research Gaps Section
        gaps = review_data.get("research_gaps", [])
        if gaps:
            gh = doc.add_heading("Identified Research Gaps and Future Trajectories", level=1)
            gh.paragraph_format.space_before = Pt(16)
            gh.paragraph_format.space_after = Pt(6)

            for g_idx, gap in enumerate(gaps, start=1):
                gp = doc.add_paragraph()
                gp.paragraph_format.space_after = Pt(6)
                g_title_run = gp.add_run(f"Gap {g_idx}: {gap.get('gap_title', '')}\n")
                g_title_run.font.name = "Times New Roman"
                g_title_run.font.size = Pt(11)
                g_title_run.font.bold = True

                desc_run = gp.add_run(f"Description: {gap.get('description', '')}\n")
                desc_run.font.name = "Times New Roman"
                desc_run.font.size = Pt(11)

                evid_run = gp.add_run(f"Grounding Evidence: {gap.get('supporting_evidence', '')}")
                evid_run.font.name = "Times New Roman"
                evid_run.font.size = Pt(10)
                evid_run.font.italic = True
                evid_run.font.color.rgb = RGBColor(71, 85, 105)

        # References Section
        ref_h = doc.add_heading("References", level=1)
        ref_h.paragraph_format.space_before = Pt(16)
        ref_h.paragraph_format.space_after = Pt(8)

        for ref_str in review_data.get("references", []):
            rp = doc.add_paragraph()
            rp.paragraph_format.left_indent = Inches(0.5)
            rp.paragraph_format.first_line_indent = Inches(-0.5)  # Hanging indent
            rp.paragraph_format.space_after = Pt(6)
            r_run = rp.add_run(ref_str)
            r_run.font.name = "Times New Roman"
            r_run.font.size = Pt(11)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path

    @staticmethod
    def generate_audit_report_docx(project_title: str, issues: List[Dict[str, Any]], stats: Dict[str, Any], output_path: str) -> str:
        """
        Creates an audit report DOCX summarizing citation findings.
        """
        doc = Document()
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run(f"Citation Audit Report: {project_title}")
        t_run.font.name = "Times New Roman"
        t_run.font.size = Pt(16)
        t_run.font.bold = True

        stat_p = doc.add_paragraph()
        stat_p.add_run(
            f"Total Issues Detected: {stats.get('total', 0)} | "
            f"Errors: {stats.get('errors', 0)} | "
            f"Warnings: {stats.get('warnings', 0)} | "
            f"Audit Status: {stats.get('status', 'Completed')}"
        )
        stat_p.paragraph_format.space_after = Pt(12)

        for idx, issue in enumerate(issues, start=1):
            h = doc.add_heading(f"Issue {idx}: [{issue.get('category')}] - {issue.get('cited_text', '')[:60]}", level=2)
            h.paragraph_format.space_before = Pt(10)

            p = doc.add_paragraph()
            p.add_run(f"Location: Page {issue.get('page')}, Paragraph {issue.get('paragraph')}\n").bold = True
            p.add_run(f"Severity: {issue.get('severity')}\n")
            p.add_run(f"Explanation: {issue.get('explanation')}\n")
            p.add_run(f"Suggested Correction: {issue.get('suggested_correction')}\n")

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path

    @staticmethod
    def generate_bibtex(sources: List[Dict[str, Any]]) -> str:
        """
        Generates standard BibTeX entries from scholarly sources.
        """
        entries = []
        for idx, src in enumerate(sources, start=1):
            authors = src.get("authors", "Author")
            first_surname = authors.split(",")[0].split()[-1].lower() if authors else "author"
            first_surname = re.sub(r'[^a-z]', '', first_surname) or "ref"
            year = src.get("year", 2023)
            cite_key = f"{first_surname}{year}_{idx}"

            entry = (
                f"@article{{{cite_key},\n"
                f"  title = {{{src.get('title', 'Unknown Title')}}},\n"
                f"  author = {{{src.get('authors', 'Unknown')}}},\n"
                f"  journal = {{{src.get('journal', 'Academic Venue')}}},\n"
                f"  year = {{{year}}},\n"
                f"  doi = {{{src.get('doi', '')}}},\n"
                f"  url = {{{src.get('url', '')}}}\n"
                f"}}"
            )
            entries.append(entry)
        return "\n\n".join(entries)
